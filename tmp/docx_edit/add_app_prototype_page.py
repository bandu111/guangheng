from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\bandu\Desktop\黑客松比赛材料\光衡_GuangHeng_Anker黑客松预赛提交材料_20页优化版(1).docx")
OUTPUT = Path(r"C:\Users\bandu\Desktop\黑客松比赛材料\光衡_GuangHeng_Anker黑客松预赛提交材料_21页_APP原型增强版.docx")

IMAGES = {
    "首页": Path(r"C:\Users\bandu\AppData\Local\Temp\codex-clipboard-7dfc8fdd-770c-4473-a9cb-5a640e49933d.png"),
    "策略": Path(r"C:\Users\bandu\AppData\Local\Temp\codex-clipboard-dc076a2a-e379-4772-98f8-c355b84f7588.png"),
    "报表": Path(r"C:\Users\bandu\AppData\Local\Temp\codex-clipboard-2c8f371f-9a2a-4cc2-aa2b-4b7ac84adc96.png"),
    "设备": Path(r"C:\Users\bandu\AppData\Local\Temp\codex-clipboard-14538564-7aa3-4df4-bdb4-3a0999bd36c9.png"),
}

BLUE = "2468D4"
NAVY = RGBColor(8, 34, 64)
GREEN = RGBColor(24, 157, 105)
MUTED = RGBColor(100, 119, 145)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7E5F7", size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def shade_cell(cell, fill="F7FAFE"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_layout_fixed(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_text(paragraph, text, size, color, bold=False, font="Microsoft YaHei"):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), font)
    return run


def set_picture_alt_text(run, title, description):
    drawings = run._element.xpath(".//wp:docPr")
    if drawings:
        drawings[0].set("name", title)
        drawings[0].set("title", title)
        drawings[0].set("descr", description)


def build_card(cell, index, title, image_path, subtitle):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, top=90, start=100, bottom=100, end=100)
    shade_cell(cell)

    label = cell.paragraphs[0]
    label.alignment = WD_ALIGN_PARAGRAPH.LEFT
    label.paragraph_format.space_before = Pt(0)
    label.paragraph_format.space_after = Pt(4)
    add_text(label, f"{index:02d}  {title}", 13, NAVY, bold=True)

    image_p = cell.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_before = Pt(0)
    image_p.paragraph_format.space_after = Pt(3)
    pic_run = image_p.add_run()
    pic_run.add_picture(str(image_path), width=Inches(3.28))
    set_picture_alt_text(pic_run, f"光衡 APP {title}原型", f"光衡 Flutter APP 的{title}页面流程原型图")

    caption = cell.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(0)
    add_text(caption, subtitle, 8.5, MUTED)


def main():
    for path in [SOURCE, *IMAGES.values()]:
        if not path.exists():
            raise FileNotFoundError(path)

    doc = Document(str(SOURCE))
    anchor = next((p for p in doc.paragraphs if p.text.strip().startswith("06 ")), None)
    if anchor is None:
        raise RuntimeError("未找到‘06 团队与后续’插入锚点")

    body = doc._element.body
    sect_pr = body.sectPr
    created = []

    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_after = Pt(5)
    add_text(eyebrow, "系统架构 + 工程证据 · 产品界面", 10.5, GREEN, bold=True)
    created.append(eyebrow._p)

    title = doc.add_paragraph(style="GH Page Title")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    add_text(title, "APP 原型总览", 25, NAVY, bold=True)
    created.append(title._p)

    rule = doc.add_table(rows=1, cols=1)
    rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    rule.autofit = False
    rule.columns[0].width = Inches(7.15)
    rule.cell(0, 0).width = Inches(7.15)
    set_cell_margins(rule.cell(0, 0), top=0, start=0, bottom=0, end=0)
    shade_cell(rule.cell(0, 0), BLUE)
    rule.cell(0, 0).paragraphs[0].paragraph_format.line_spacing = Pt(2)
    rule.cell(0, 0).paragraphs[0].paragraph_format.space_after = Pt(0)
    created.append(rule._tbl)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(7)
    intro.paragraph_format.space_after = Pt(8)
    add_text(intro, "四个主入口覆盖家庭能源总览、策略决策、价值报表与设备管理，共享同一套真实状态、权限、安全、执行与验证闭环。", 9.5, MUTED)
    created.append(intro._p)

    grid = doc.add_table(rows=2, cols=2)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    set_repeat_table_layout_fixed(grid)
    set_table_borders(grid)
    for row in grid.rows:
        for cell in row.cells:
            cell.width = Inches(3.55)

    build_card(grid.cell(0, 0), 1, "首页", IMAGES["首页"], "家庭能源流向、关键指标、主动建议与最近一次验证。")
    build_card(grid.cell(0, 1), 2, "策略", IMAGES["策略"], "24 小时计划、策略偏好、跨设备协同与一次确认。")
    build_card(grid.cell(1, 0), 3, "报表", IMAGES["报表"], "今日 / 本周 / 本月趋势、决策执行记录与日历支出。")
    build_card(grid.cell(1, 1), 4, "设备", IMAGES["设备"], "设备发现、健康诊断、家庭负载、实时数据与控制能力。")
    created.append(grid._tbl)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(7)
    note.paragraph_format.space_after = Pt(0)
    add_text(note, "真实 Flutter 原型 · 首页 / 策略 / 报表 / 设备", 9, GREEN, bold=True)
    created.append(note._p)

    end_break = doc.add_paragraph()
    end_break.add_run().add_break()
    end_break.runs[0]._element.br_lst[0].set(qn("w:type"), "page")
    created.append(end_break._p)

    anchor_element = anchor._p
    for element in created:
        if element.getparent() is not None:
            element.getparent().remove(element)
        anchor_element.addprevious(element)

    # Keep the section properties as the final body child.
    if sect_pr is not None and sect_pr.getparent() is body:
        body.remove(sect_pr)
        body.append(sect_pr)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
