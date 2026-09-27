from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

ROOT = Path(r"C:\Users\bandu\Desktop\黑客松比赛材料")
IMG = ROOT / "assets" / "generated"
OUT = ROOT / "光衡_GuangHeng_Anker黑客松预赛提交材料.docx"

NAVY = "0B1F3A"; BLUE="2478F3"; GREEN="16B77A"; ORANGE="F5A623"; INK="17263D"; MUTED="66758B"; PALE="F4F7FB"; LINE="DDE6F1"; RED="E65B55"

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.52); sec.bottom_margin=Inches(.48); sec.left_margin=Inches(.58); sec.right_margin=Inches(.58)

styles=doc.styles
styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(9.2); styles['Normal'].font.color.rgb=RGBColor.from_string(INK)
styles['Title'].font.name='Microsoft YaHei'; styles['Title']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Title'].font.size=Pt(34); styles['Title'].font.bold=True; styles['Title'].font.color.rgb=RGBColor.from_string(NAVY)
for name,size,color in [('Heading 1',24,NAVY),('Heading 2',16,BLUE),('Heading 3',12,INK)]:
    st=styles[name]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(6); st.paragraph_format.space_after=Pt(5)

if 'Takeaway' not in styles:
    st=styles.add_style('Takeaway',WD_STYLE_TYPE.PARAGRAPH)
    st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(13); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(NAVY)
    st.paragraph_format.space_before=Pt(5); st.paragraph_format.space_after=Pt(7)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def margins(cell,top=70,start=90,bottom=70,end=90):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def table(rows,widths=None,header=True,font_size=8.4):
    t=doc.add_table(rows=len(rows), cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            c=t.cell(i,j); c.text=str(val); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
            if widths: c.width=Inches(widths[j])
            if i==0 and header:
                shade(c,NAVY)
                for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(font_size)
            else:
                if i%2==0: shade(c,PALE)
                for r in c.paragraphs[0].runs: r.font.size=Pt(font_size); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        if i==0 and header: set_repeat_table_header(t.rows[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def p(text='',style=None,size=None,color=None,bold=False,align=None,space_after=4):
    q=doc.add_paragraph(style=style)
    q.paragraph_format.space_after=Pt(space_after)
    if align is not None: q.alignment=align
    r=q.add_run(text); r.bold=bold
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    return q

def bullets(items,size=9.2):
    for item in items:
        q=doc.add_paragraph(style='List Bullet'); q.paragraph_format.space_after=Pt(2); q.paragraph_format.left_indent=Inches(.2)
        r=q.add_run(item); r.font.size=Pt(size); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')

def img(name,width=7.25,caption=None):
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(3)
    q.add_run().add_picture(str(IMG/name),width=Inches(width))
    if caption: p(caption,size=7.5,color=MUTED,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=2)

def pagebreak(): doc.add_page_break()

def page_label(text):
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.RIGHT; q.paragraph_format.space_after=Pt(1)
    r=q.add_run(text.upper()); r.font.size=Pt(7.5); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(GREEN)

# Footer page number
for section in doc.sections:
    f=section.footer; q=f.paragraphs[0]; q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=q.add_run('GUANGHENG × ANKER HACKATHON  ·  '); r.font.size=Pt(7); r.font.color.rgb=RGBColor.from_string(MUTED)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); q._p.append(fld)

# P1 cover
p('ANKER 首届黑客松挑战赛 · 充电储能赛道',size=10,color=GREEN,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=24)
p('光衡',style='Title',align=WD_ALIGN_PARAGRAPH.CENTER,space_after=0)
p('GuangHeng',size=22,color=BLUE,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=8)
p('AUTONOMOUS HOME ENERGY AGENT',size=13,color=MUTED,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,space_after=18)
p('家庭能源管理真正缺少的已经不是更多数据，\n而是从数据到行动的持续决策能力。',style='Takeaway',align=WD_ALIGN_PARAGRAPH.CENTER,space_after=18)
img('01_problem_before_after.png',7.0)
table([
    ['家庭能源 Autopilot','受治理的 AI 安全执行','一个 Agent · 两种入口'],
    ['Observe · Predict · Optimize','Approval · Safety · Verify','Flutter + ESP32-S3']
],widths=[2.35,2.35,2.35],font_size=8.5)
p('当前工程验证版本｜SOLIX 设备闭环：SIMULATOR_E2E｜ESP32-S3：REAL_HIL',size=8,color=MUTED,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
pagebreak()

# P2 info + challenge
page_label('01 · PRELIMINARY INFORMATION')
doc.add_heading('一、预赛信息卡',level=1)
table([
    ['项目','内容'],['队名','【待填写】'],['意向赛道','充电储能'],
    ['一句话摘要','运行在 Home Assistant 之上的自主家庭能源 Agent：持续预测与优化，只在需要时请求授权，并用回读和 Smart Meter 证明结果。'],
    ['队长','【待填写：姓名 / 邮箱 / 手机号】'],['成员介绍 & 分工','【待填写】']
],widths=[1.45,5.7],font_size=8.2)
doc.add_heading('二、方案成果展示',level=1)
doc.add_heading('1. 赛题理解',level=2)
p('核心矛盾：Decision Burden，而不是 Data Visibility。',style='Takeaway')
p('很多系统已经能展示 PV、负载、SOC 与电网功率；用户仍需每天判断何时充放电、留多少备电、天气如何影响明天，以及多个设备如何协同。')
bullets(['数据多但行动门槛高：普通家庭被迫理解复杂曲线与设备边界。','控制响应式且单设备化：缺少预测和跨设备计划。','命令成功不等于家庭目标达成：缺少写后回读和结果证明。'])
p('目标用户：拥有或计划拥有光伏、储能、智能电表和可控负载，希望提高光伏自用、合理使用储能并降低重复操作的家庭。',bold=True,color=NAVY)
pagebreak()

# P3 core 1
page_label('02 · CORE FEATURE 1')
doc.add_heading('2. 具体方案说明（突出 AI 能力）',level=1)
doc.add_heading('核心功能一｜家庭能源 Autopilot',level=2)
p('系统持续观察和判断；No Action 时保持安静，需要行动时才生成 Proposal。',style='Takeaway')
img('03_backend_business_loop.png',7.2)
table([
    ['环节','实现'],['Observe','Home Assistant 状态、历史、设备能力与 Smart Meter'],
    ['Predict','Weather、Solar Forecast、Load Forecast、参考电价'],
    ['Optimize','Deterministic Optimizer + SAVE / AUTO / BACKUP 硬约束'],
    ['Decide','已满足则不打扰；需要行动则创建 Proposal / Action Set']
],widths=[1.3,5.8],font_size=8.5)
p('AI 的角色',style='Heading 3')
p('Hermes 理解用户目标、进行 What-if、解释“为什么现在”并生成结构化意图；Reserve、功率等数值由确定性 Optimizer 计算，避免 LLM 发明目标。')
pagebreak()

# P4 core2
page_label('03 · CORE FEATURE 2')
doc.add_heading('核心功能二｜受治理的 AI 安全执行闭环',level=2)
p('AI can reason. AI can propose. AI cannot freely execute.',style='Takeaway')
img('04_ai_safety_boundary.png',7.2)
bullets(['14 个 MCP Tools 仅允许查询、评估与提案；无 approve、execute、HA write。','执行前重新读取 Runtime，并检查 control enabled、device online、capability、权限、verified、范围、步长与审批。','Action Set 顺序执行、逐项回读、失败即停；未执行动作标记 SKIPPED。','只有 SUCCEEDED + VERIFIED 才显示“已验证”。'])
table([
    ['验证层级','证明什么'],['L1 · Command Accepted','请求进入系统'],['L2 · Device Readback','设备状态与目标一致'],['L3 · Household Outcome','Smart Meter 证明家庭能源目标实现']
],widths=[2.1,5.0],font_size=8.6)
pagebreak()

# P5 core3
page_label('04 · CORE FEATURE 3')
doc.add_heading('核心功能三｜一个 Agent，两种产品入口',level=2)
p('Flutter 提供深度管理；ESP32-S3 提供 Ambient Physical Interaction；两端共享同一 Backend Truth。',style='Takeaway')
img('08_flutter_product_evidence.png',7.1,'真实 Flutter 页面：策略、Action Set 与参考基线报表')
table([
    ['入口','产品职责'],['Flutter','主要用户产品端：首页、策略、报表、设备、审批、执行、验证与系统健康'],
    ['ESP32-S3 Companion','家庭空间中的 Ambient 状态、配对、解释与 nonce/version 绑定的物理长按确认'],
    ['Shared Truth','Proposal、Approval、Execution 与 Verification 统一由 Backend / Control Plane 管理']
],widths=[1.8,5.3],font_size=8.1)
pagebreak()

# P6 value/mvp
page_label('05 · VALUE & MVP')
doc.add_heading('3. 方案价值与预期效果',level=2)
table([
    ['价值','预期改善'],['降低认知负担','从“读懂曲线再操作”变成系统持续观察、必要时才打扰'],
    ['预测式能源管理','天气、光伏与负载预测进入当前决策'],['多设备协同','Storage + Meter + Flexible Load 形成 Action Set'],
    ['结果可证明','设备回读 + Smart Meter before/after，而非把 HTTP 200 当成功']
],widths=[1.8,5.3],font_size=8.4)
doc.add_heading('24 小时可演示 MVP',level=2)
p('核心是完整业务闭环，不是堆叠页面，也不是主要造硬件。',style='Takeaway')
img('07_end_to_end_flow.png',7.15)
p('ESP32-S3 Companion 作为加分实体入口，演示配对、Ambient 状态和物理长按；真实 Anker 家庭硬件与长期节省效果留到赛后实机阶段。',size=8.7,color=MUTED)
doc.add_heading('4. 方案优势 / 创新点',level=2)
bullets(['Dashboard → Autonomous Home Energy Agent','Deterministic Optimization + Agent Reasoning','AI Intelligence 与 Execution Authority 分离','Cross-device Action Set + Fail-stop','L1 / L2 / L3 Unified Verification'],size=8.6)
pagebreak()

# P7 architecture
page_label('06 · FREE DISPLAY / ARCHITECTURE')
doc.add_heading('三、自由展示区',level=1)
doc.add_heading('完整产品架构｜Control Plane 是业务真相',level=2)
img('02_full_product_architecture.png',7.2)
table([
    ['层','真实职责'],['GuangHeng Backend','Control Plane：业务真相、决策、权限、执行与验证'],['Home Assistant','Device Integration Layer：Entity、History、Service Control'],['Hermes','Reasoning / Intent / Explanation；无自由执行权'],['Optimizer','可复现数值目标与硬约束'],['Companion','Ambient Interface；不是第二套执行链']
],widths=[1.8,5.3],font_size=8.2)
pagebreak()

# P8 real evidence
page_label('07 · FREE DISPLAY / REAL PRODUCT')
doc.add_heading('真实产品｜Agent 进入家庭空间',level=2)
p('ESP32-S3 本体能力按 REAL_HIL 记录；Anker 能源设备闭环按 SIMULATOR_E2E 记录，两者不混写。',style='Takeaway')
img('09_companion_hil_evidence.png',7.2)
table([
    ['证据','状态'],['ESP32-S3 rev 0.2 / 16 MB Flash / 8 MB PSRAM','REAL_HIL'],['AMOLED / Touch / Wi-Fi / SNTP / TLS / Pairing','REAL_HIL'],['物理 1.4 s 长按审批','REAL_HIL'],['SOLIX 多设备写入 / 回读 / Smart Meter L3','SIMULATOR_E2E'],['真实 Anker 家庭硬件','NOT_VERIFIED']
],widths=[5.2,1.9],font_size=8.4)
pagebreak()

# P9 evidence/team
page_label('08 · EVIDENCE & ROADMAP')
doc.add_heading('工程证据与真实性边界',level=2)
table([
    ['分类','当前证据'],['REAL_HIL','ESP32-S3 实板、AMOLED / Touch、Wi-Fi / SNTP / TLS、配对、Snapshot、物理长按'],
    ['SIMULATOR_E2E','SOLIX Capability、Action Set、Safety、顺序执行、设备回读、Smart Meter L3'],
    ['SOFTWARE_IMPLEMENTED','Backend 1.6.0；Flutter 1.4.0+9；14 MCP Tools；Autonomy / Optimizer / Voice 等模块'],
    ['PARTIAL / HIL_PENDING','Lift-to-Explain 校准；ESP32-S3 Voice 全链路；Verification 最终视觉验收'],
    ['PLANNED P2','TMAG5273 / Magnetic Context'],['NOT_VERIFIED','真实 Anker 家庭硬件、长期节省、准确率、用户规模']
],widths=[1.65,5.45],font_size=7.9)
p('当前回归证据（2026-09-26）',style='Heading 3')
p('Backend：156 passed｜Flutter：15 passed｜Firmware host tests 与 ESP-IDF esp32s3 build 通过。测试数量仅作为本次工程审计记录。',size=8.5)
doc.add_heading('团队与后续',level=2)
table([['队名','【待填写】'],['队长','【待填写：姓名 / 邮箱 / 手机号】'],['成员与分工','【待填写】']],widths=[1.5,5.6],header=False,font_size=8.8)
bullets(['P0：真实 Anker SOLIX 硬件逐项完成写入、回读与 L3 场景验证。','P1：完成 ESP32-S3 Voice HIL 与 Lift-to-Explain 校准。','P2：再引入 TMAG5273 / Magnetic Context，不提前包装为当前能力。'],size=8.6)
p('真实性声明',style='Heading 3')
p('本材料未使用虚构用户数、准确率、节省比例或 Anker 实机结果。报表费用为参考基线估算；当前 Anker 多设备闭环属于 Simulator E2E。',size=8.3,color=RED,bold=True)

doc.core_properties.title='光衡 GuangHeng｜Anker 黑客松预赛提交材料'
doc.core_properties.subject='充电储能赛道｜Autonomous Home Energy Agent'
doc.core_properties.author='【待填写】'
doc.save(OUT)
print(OUT)
