from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\bandu\Desktop\黑客松比赛材料\光衡_GuangHeng_Anker黑客松预赛提交材料_20页优化版(1).docx")
OUTPUT = Path(r"C:\Users\bandu\Desktop\黑客松比赛材料\光衡_GuangHeng_Anker黑客松预赛提交材料_最终强化版.docx")

NAVY = "0B1F3A"
BLUE = "235FB3"
GREEN = "00A87A"
PALE_BLUE = "EAF2FB"
PALE_GREEN = "EAF7F1"
PALE_YELLOW = "FFF5DE"
PALE_RED = "FDECEC"
MID_GRAY = "5C6878"
GRID = "D7DEE8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def unique_cells(row):
    result = []
    seen = set()
    for cell in row.cells:
        marker = id(cell._tc)
        if marker not in seen:
            seen.add(marker)
            result.append(cell)
    return result


def format_run(run, *, size=8.5, bold=False, color=NAVY, font="Aptos") -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")


def set_paragraph_text(paragraph, text: str, *, size=None, bold=None, color=None, align=None) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    if any(v is not None for v in (size, bold, color)):
        format_run(
            run,
            size=size if size is not None else 10.5,
            bold=bool(bold),
            color=color or NAVY,
        )
    if align is not None:
        paragraph.alignment = align


def set_cell_text(
    cell,
    text: str,
    *,
    size=8.2,
    bold=False,
    color=NAVY,
    fill=None,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    format_run(run, size=size, bold=bold, color=color)
    if fill:
        set_cell_shading(cell, fill)
    set_cell_border(cell)
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side in ("top", "left", "bottom", "right"):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), "90")
        node.set(qn("w:type"), "dxa")


def find_paragraph(document: Document, exact: str):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise KeyError(f"Paragraph not found: {exact}")


def find_row(table, key: str):
    for row in table.rows:
        if any(key in cell.text for cell in unique_cells(row)):
            return row
    raise KeyError(f"Row not found: {key}")


def replace_visual_row(table, key: str, values: list[str], *, size=8.0) -> None:
    row = find_row(table, key)
    cells = unique_cells(row)
    if len(cells) == 1 and len(values) > 1:
        set_cell_text(
            cells[0],
            "\n".join(values),
            size=size,
            bold=False,
            color=NAVY,
            fill=PALE_YELLOW,
        )
        return
    if len(cells) < len(values):
        raise ValueError(f"Row '{key}' has {len(cells)} visual cells, need {len(values)}")
    for index, value in enumerate(values):
        set_cell_text(
            cells[index],
            value,
            size=size,
            bold=index == 0,
            color=NAVY if index != len(values) - 1 else GREEN,
            fill=PALE_BLUE if index == 0 else WHITE,
        )


def style_table(table, widths=None, font_size=8.2, header=True) -> None:
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for col_index, cell in enumerate(unique_cells(row)):
            if widths and col_index < len(widths):
                cell.width = Inches(widths[col_index])
            text = cell.text
            fill = PALE_BLUE if header and row_index == 0 else (PALE_GREEN if row_index % 2 == 0 else WHITE)
            set_cell_text(
                cell,
                text,
                size=font_size,
                bold=header and row_index == 0,
                color=NAVY,
                fill=fill,
            )


def add_eyebrow(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    format_run(run, size=8, bold=True, color=GREEN)
    return paragraph


def add_divider(document: Document, template_table):
    element = deepcopy(template_table._tbl)
    document.element.body.insert(-1, element)
    return element


def move_elements_before(anchor, elements) -> None:
    for element in elements:
        anchor.addprevious(element)


def add_ai_page(document: Document, anchor, divider_template) -> None:
    elements = []
    eyebrow = add_eyebrow(document, "AI TECHNICAL PATH")
    elements.append(eyebrow._p)
    title = document.add_paragraph(style="GH Page Title")
    set_paragraph_text(title, "AI 技术路径：让 AI 负责理解，让确定性系统负责控制")
    elements.append(title._p)
    elements.append(add_divider(document, divider_template))
    subtitle = document.add_paragraph(style="GH Quote")
    set_paragraph_text(
        subtitle,
        "光衡不是让 LLM 决定设备功率，而是把自然语言的不确定性和真实设备控制的确定性分开处理。",
        size=9.5,
        bold=True,
        color=NAVY,
    )
    subtitle.paragraph_format.space_after = Pt(7)
    elements.append(subtitle._p)

    section = document.add_paragraph(style="GH Section")
    set_paragraph_text(section, "模块职责与设计理由")
    elements.append(section._p)

    rows = [
        ("模块", "职责", "为什么这样设计"),
        ("Hermes AI", "理解自然语言；识别家庭能源目标；解释方案原因；回答 What-if；目标不明确时继续澄清。", "自然语言和用户目标存在不确定性，适合 AI 处理。"),
        ("Forecast Layer", "提供天气、未来光伏和未来负载趋势，为当前判断补充未来上下文。", "预测只提供决策上下文，不能因此获得设备控制权。"),
        ("Deterministic Optimizer", "根据 SOC、用户策略、设备能力、目标范围与硬约束，计算具体目标值。", "相同输入得到可重复结果，避免 LLM 凭感觉生成控制参数。"),
        ("执行前安全检查", "执行前重新检查用户权限、设备在线、数据有效性、设备能力、安全范围与最新状态。", "授权后状态仍可能变化，执行前必须重新确认。"),
        ("Execution", "通过 Home Assistant / SOLIX 能力按一次目标下的多设备行动方案顺序执行。", "所有真实写入必须经过受治理的系统控制中枢。"),
        ("Verification", "L1 命令确认；L2 设备确认；L3 Smart Meter 家庭结果确认。", "AI 的解释不是成功证据，最终以设备与家庭能源数据为准。"),
    ]
    table = document.add_table(rows=len(rows), cols=3)
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            set_cell_text(
                table.cell(r, c),
                value,
                size=7.7 if r else 8.1,
                bold=r == 0 or c == 0,
                fill=PALE_BLUE if r == 0 else (PALE_GREEN if c == 0 else WHITE),
            )
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(unique_cells(row), (1.25, 3.55, 2.35)):
            cell.width = Inches(width)
    elements.append(table._tbl)

    reason = document.add_paragraph(style="GH Section")
    set_paragraph_text(reason, "为什么不让 LLM 直接决定设备功率？")
    reason.paragraph_format.space_before = Pt(7)
    reason.paragraph_format.space_after = Pt(3)
    elements.append(reason._p)
    callout = document.add_table(rows=2, cols=1)
    set_cell_text(
        callout.cell(0, 0),
        "家庭储能是真实物理系统。语言模型适合理解“今晚多留一点电”“明天下雨提前准备”等模糊目标；具体充放电目标、设备边界与执行顺序必须可重复、可测试、可约束。",
        size=8.0,
        fill=PALE_YELLOW,
    )
    set_cell_text(
        callout.cell(1, 0),
        "AI 理解  +  确定性计算  +  受治理执行\nAI 的价值不是替代所有算法，而是把人的自然语言目标转换为系统能够安全处理的能源任务。",
        size=8.2,
        bold=True,
        color=GREEN,
        fill=PALE_GREEN,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    elements.append(callout._tbl)
    page_break = document.add_page_break()
    elements.append(page_break._p)
    move_elements_before(anchor, elements)


def add_value_solix_page(document: Document, anchor, divider_template) -> None:
    elements = []
    eyebrow = add_eyebrow(document, "VALUE MODEL + ANKER SOLIX")
    elements.append(eyebrow._p)
    title = document.add_paragraph(style="GH Page Title")
    set_paragraph_text(title, "价值如何量化，以及光衡对 Anker SOLIX 的价值")
    elements.append(title._p)
    elements.append(add_divider(document, divider_template))
    intro = document.add_paragraph()
    set_paragraph_text(
        intro,
        "光衡当前能够验证单次行动是否有效；长期运行后，同一套数据链可以进一步计算真实家庭价值。",
        size=9.3,
        bold=True,
        color=NAVY,
    )
    intro.paragraph_format.space_after = Pt(6)
    elements.append(intro._p)

    formula = document.add_table(rows=3, cols=1)
    set_cell_text(formula.cell(0, 0), "价值量化模型", size=9, bold=True, fill=PALE_BLUE, color=BLUE)
    set_cell_text(
        formula.cell(1, 0),
        "新增自用电量 ＝ 年光伏发电量 ×（优化后自用率 − Baseline 自用率）\n\n年净收益 ≈ 新增自用电量 ×（购电价 − 上网电价）＋ 峰谷调度收益 − 电池额外循环成本",
        size=9.0,
        bold=True,
        fill=WHITE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell_text(
        formula.cell(2, 0),
        "示例测算｜仅用于说明模型，不代表当前实测收益\n光伏平均发电 10 kWh/日；自用率 30% → 45%，提升 15 个百分点；年光伏发电约 3650 kWh；额外家庭自用约 548 kWh/年。最终经济收益仍取决于当地购电价、上网电价、峰谷价差和电池循环成本。预赛阶段不把该示例作为真实长期节省结论。",
        size=7.8,
        fill=PALE_YELLOW,
    )
    elements.append(formula._tbl)

    section = document.add_paragraph(style="GH Section")
    set_paragraph_text(section, "对 Anker SOLIX 的价值")
    section.paragraph_format.space_before = Pt(7)
    section.paragraph_format.space_after = Pt(3)
    elements.append(section._p)
    solix_rows = [
        ("01｜从单设备到家庭级协同", "Solarbank、Smart Meter 与柔性负载不再是彼此独立的入口。光衡围绕家庭目标组织设备能力，形成一次目标下的多设备行动方案。"),
        ("02｜Smart Meter 成为价值证明设备", "Smart Meter 不只展示功率，而是承担 L3 家庭结果确认，让“减少反送、降低购电、改变峰值”变成用户可感知的证据。"),
        ("03｜提升 SOLIX 生态持续使用价值", "用户持续使用同一套目标、策略、执行和验证闭环；新设备可按设备能力进入同一个 Agent，无需重建体验。"),
    ]
    solix = document.add_table(rows=3, cols=2)
    for r, (heading, body) in enumerate(solix_rows):
        set_cell_text(solix.cell(r, 0), heading, size=8.1, bold=True, fill=PALE_BLUE, color=BLUE)
        set_cell_text(solix.cell(r, 1), body, size=7.8, fill=WHITE)
        solix.cell(r, 0).width = Inches(2.2)
        solix.cell(r, 1).width = Inches(4.95)
    elements.append(solix._tbl)
    close = document.add_paragraph(style="GH Quote")
    set_paragraph_text(
        close,
        "对 Anker，光衡的价值不是增加一个新的 Dashboard，而是让 Solarbank、Smart Meter 与可控设备在同一个家庭目标下形成持续协同。",
        size=8.7,
        bold=True,
        color=GREEN,
    )
    close.paragraph_format.space_before = Pt(5)
    elements.append(close._p)
    page_break = document.add_page_break()
    elements.append(page_break._p)
    move_elements_before(anchor, elements)


def add_evidence_links(document: Document, anchor) -> None:
    elements = []
    table = document.add_table(rows=5, cols=2)
    header = table.cell(0, 0).merge(table.cell(0, 1))
    set_cell_text(header, "可验证工程证据", size=8.5, bold=True, fill=PALE_BLUE, color=BLUE)
    rows = [
        ("完整 REAL_E2E 演示", "[二维码 / 视频链接待补]"),
        ("工程仓库 / Evidence Repo", "[链接待补]"),
        ("CI / Test Report", "[链接待补]"),
        ("覆盖范围", "发现反送 → 生成行动方案 → 用户授权 → 执行 → 设备状态回读 → Smart Meter L3"),
    ]
    for index, (left, right) in enumerate(rows, start=1):
        set_cell_text(table.cell(index, 0), left, size=7.1, bold=True, fill=PALE_GREEN)
        set_cell_text(table.cell(index, 1), right, size=7.1, fill=WHITE)
    elements.append(table._tbl)
    move_elements_before(anchor, elements)


def add_demo_fallback(document: Document, anchor) -> None:
    elements = []
    table = document.add_table(rows=5, cols=2)
    header = table.cell(0, 0).merge(table.cell(0, 1))
    set_cell_text(header, "Demo Reliability｜现场容灾", size=8.4, bold=True, fill=PALE_BLUE, color=BLUE)
    rows = [
        ("首选链路", "真实 Anker SOLIX + Home Assistant + Smart Meter（现场实机条件就绪时）"),
        ("Fallback A｜网络 / 云端异常", "停止新的真实设备写入，不使用过期状态继续执行。"),
        ("Fallback B｜Simulator 回放", "重放同一 Action Set、执行前安全检查、Fail-stop 与 L3 验证链路。"),
        ("Fallback C｜真实闭环证据", "[完整 REAL_E2E 演示二维码 / 视频链接待补]"),
    ]
    for index, (left, right) in enumerate(rows, start=1):
        set_cell_text(table.cell(index, 0), left, size=6.8, bold=True, fill=PALE_GREEN)
        set_cell_text(table.cell(index, 1), right, size=6.8, fill=WHITE)
    elements.append(table._tbl)
    note = document.add_paragraph(style="GH Small")
    set_paragraph_text(
        note,
        "当前可审计能源设备闭环证据为 SOLIX-compatible Simulator E2E；真实 Anker 家庭硬件仍待实机验证。",
        size=7.0,
        color=MID_GRAY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    note.paragraph_format.space_before = Pt(2)
    elements.append(note._p)
    move_elements_before(anchor, elements)


def main() -> None:
    document = Document(SOURCE)
    tables = list(document.tables)
    divider_template = tables[1]

    # Cover: make L3 the first remembered differentiator.
    set_paragraph_text(
        document.paragraphs[3],
        "光衡是一套能够持续观察、主动判断、协调多设备执行，并用 Smart Meter 证明家庭能源目标是否真正实现的 Autonomous Home Energy Agent。系统平时保持安静，只在值得行动时生成方案；用户授权后进入受治理的安全执行链，最终不仅确认“命令发出”和“设备改变”，还继续验证家庭层面的真实结果。",
        size=9.7,
        color=NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    cover = tables[0]
    eq_cell = cover.cell(0, 0).merge(cover.cell(0, 1))
    set_cell_text(
        eq_cell,
        "命令成功  ≠  设备成功  ≠  家庭目标成功",
        size=13.2,
        bold=True,
        color=BLUE,
        fill=PALE_BLUE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell_text(
        cover.cell(1, 0),
        "L1 命令确认\n→ L2 设备确认\n→ L3 Smart Meter 家庭结果确认",
        size=9.6,
        bold=True,
        color=GREEN,
        fill=PALE_GREEN,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell_text(
        cover.cell(1, 1),
        "主动判断 + 多设备行动方案\n用户授权 + 执行前安全检查\n设备状态回读 + 家庭结果证明",
        size=8.5,
        bold=True,
        color=NAVY,
        fill=WHITE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_paragraph_text(
        document.paragraphs[5],
        "证据边界：多设备执行、设备状态回读与 Smart Meter L3 已完成 SOLIX-compatible Simulator E2E；ESP32-S3 物理长按授权已完成 REAL_HIL；真实 Anker 家庭硬件写入仍待实机验证。",
        size=7.2,
        color=MID_GRAY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    # Submission card and claim hygiene.
    profile = tables[2]
    replace_visual_row(
        profile,
        "一句话摘要",
        [
            "一句话摘要",
            "光衡持续观察并预测家庭能源状态，在真正值得行动时主动生成多设备方案；用户授权后由安全执行链完成设备控制，并通过设备状态回读与 Smart Meter 三级验证，证明家庭能源目标是否真正达成。",
        ],
        size=7.8,
    )
    snapshot = tables[4]
    replace_visual_row(snapshot, "真实 Anker SOLIX", ["SOLIX-compatible 写入 / 回读", "SIMULATOR_E2E｜VERIFIED"], size=7.6)
    replace_visual_row(snapshot, "真实跨设备 Action Set", ["跨设备 Action Set（4 项）", "SIMULATOR_E2E｜VERIFIED"], size=7.6)
    replace_visual_row(snapshot, "Smart Meter 家庭结果验证", ["Smart Meter L3 家庭结果确认", "SIMULATOR_E2E｜VERIFIED"], size=7.6)
    replace_visual_row(snapshot, "Simulator 异常 / 回归闭环", ["ESP32-S3 物理长按授权", "REAL_HIL｜VERIFIED"], size=7.6)
    replace_visual_row(snapshot, "Voice：ESP32-S3", ["Voice：软件 / 服务链路", "SOFTWARE_IMPLEMENTED｜HIL_PENDING"], size=7.3)
    replace_visual_row(snapshot, "事实边界", ["事实边界", "真实 Anker 家庭硬件写入、长期节省与实板 Voice HIL 尚未验证"], size=7.2)

    pain_title = find_paragraph(document, "五个真实用户痛点：从“数据展示”走向“主动闭环”")
    set_paragraph_text(pain_title, "五个核心用户痛点：从“数据展示”走向“主动闭环”")

    # Insert the AI technical path after the Autopilot capability page.
    ai_anchor = find_paragraph(document, "真实 FLUTTER 页面")._p
    add_ai_page(document, ai_anchor, divider_template)

    # Evidence-safe wording for Action Set and Smart Meter pages.
    evidence_summary = tables[19]
    replace_visual_row(evidence_summary, "真实 Anker SOLIX", ["SOLIX-compatible 写入 / 回读", "SIMULATOR_E2E｜VERIFIED"], size=7.8)
    replace_visual_row(evidence_summary, "真实跨设备 Action Set", ["跨设备 Action Set（4 项）", "SIMULATOR_E2E｜VERIFIED"], size=7.8)
    replace_visual_row(evidence_summary, "Smart Meter 家庭结果验证", ["Smart Meter L3 家庭结果确认", "SIMULATOR_E2E｜VERIFIED"], size=7.8)

    e2e_title = find_paragraph(document, "真实设备 E2E 成果：从方案到家庭结果")
    set_paragraph_text(e2e_title, "完整 Action Set E2E 成果：从方案到家庭结果")
    e2e_copy = find_paragraph(document, "这组真实产品页面对应的不是单设备开关，而是一套多设备 Action Set：一次授权后，系统按顺序执行并回读每个设备，最后以 Smart Meter 的家庭结果作为统一验证。")
    set_paragraph_text(
        e2e_copy,
        "这组真实产品页面对应一套 4 项多设备 Action Set：参与角色包括储能、Smart Meter 与可控负载 profiles；一次授权后顺序执行、逐步回读、异常即 Fail-stop，最终由 Smart Meter 统一验证家庭结果。当前能源设备链路证据来自 SOLIX-compatible Simulator E2E。",
        size=8.4,
    )
    e2e_table = tables[35]
    replace_visual_row(e2e_table, "真实 Anker SOLIX", ["SOLIX-compatible 写入 / 回读", "Home Assistant + TCP Simulator 受控写入，随后重读目标状态", "SIMULATOR_E2E｜VERIFIED"], size=6.9)
    replace_visual_row(e2e_table, "跨设备 Action Set", ["4 项跨设备 Action Set", "一次授权 → 顺序执行 → 每步状态回读 → 异常 Fail-stop", "SIMULATOR_E2E｜VERIFIED"], size=6.9)
    replace_visual_row(e2e_table, "Smart Meter 家庭结果", ["Smart Meter L3", "记录执行前后家庭电网变化，并与预期方向对比", "SIMULATOR_E2E｜VERIFIED"], size=6.9)
    replace_visual_row(e2e_table, "Simulator 的正确定位", ["ESP32-S3 物理授权", "真实实板长按批准一次 4 项 Action Set；能源设备侧仍为 Simulator", "REAL_HIL｜VERIFIED"], size=6.9)
    e2e_break = e2e_title._p
    while e2e_break is not None:
        e2e_break = e2e_break.getnext()
        if e2e_break is not None and len(e2e_break.xpath('.//w:br[@w:type="page"]')):
            break
    if e2e_break is None:
        raise RuntimeError("Could not find E2E page break")
    add_evidence_links(document, e2e_break)

    # Voice becomes an enhanced Companion capability, not the main demo.
    voice_title = find_paragraph(document, "核心功能 3｜自然交互：Flutter + Voice + Energy Companion")
    set_paragraph_text(voice_title, "核心功能 3｜一个能源大脑，两种产品入口：Flutter + Energy Companion")
    set_paragraph_text(find_paragraph(document, "用户表达目标，而不是研究设备菜单和参数。"), "同一个能源大脑，由 Flutter 深度管理，由 Energy Companion 即时陪伴。")
    voice_body = find_paragraph(document, "光衡怎么做：Flutter 负责完整管理；Energy Companion 负责即时、自然、家庭空间中的交互；Voice 让用户直接说出目标。所有入口都连接同一个 GuangHeng Backend，共享同一份状态、权限、方案和执行结果。")
    set_paragraph_text(
        voice_body,
        "光衡怎么做：Flutter 负责完整管理，Energy Companion 负责随身查看、提醒与长按授权；两端连接同一个系统控制中枢，共享状态、权限、行动方案与验证结果。Voice 是 Energy Companion 的增强交互能力之一，当前软件与服务链路已实现，ESP32-S3 真实 Voice HIL 待最终验收。",
        size=8.6,
    )
    set_paragraph_text(find_paragraph(document, "高价值语音示例"), "Energy Companion 的增强交互：Voice")
    set_paragraph_text(find_paragraph(document, "语音执行闭环"), "语音链路状态｜SOFTWARE_IMPLEMENTED · HIL_PENDING")
    voice_summary = tables[39]
    cells = unique_cells(voice_summary.rows[0])
    set_cell_text(cells[0], "当前状态", size=7.8, bold=True, fill=PALE_BLUE)
    set_cell_text(cells[-1], "软件链路已实现；ESP32-S3 真实 Voice HIL 待最终验收。主 Demo 仍以 Autopilot → Proposal → Approval → Action Set → 设备状态回读 → Smart Meter L3 为准。", size=7.5, fill=PALE_GREEN)

    # Demo storyline plus robust fallback. Avoid claiming unverified real hardware.
    demo_rows = [
        (tables[54], "01", "设备接入｜Home Assistant 接入 SOLIX-compatible 储能、Smart Meter 与可控负载 profiles。"),
        (tables[55], "02", "发现机会｜中午 Smart Meter 检测到明显光伏反送。"),
        (tables[56], "03", "主动判断｜光衡发现储能仍有可充空间，可控负载也可以运行。"),
        (tables[57], "04", "生成方案｜Flutter 显示“减少光伏反送”的多设备方案，解释原因、设备角色与预期效果。"),
        (tables[58], "05", "用户授权｜用户在 Flutter 确认，或在 Energy Companion 上长按授权。"),
        (tables[59], "06", "安全执行｜系统读取最新状态并逐项检查后，按顺序控制储能与可控负载。"),
        (tables[60], "07", "设备确认｜每个动作完成后重新读取设备状态；任何一步失败立即停止后续动作。"),
        (tables[61], "08", "家庭结果确认｜Smart Meter 再次读取电网功率；反送下降后显示“方案已完成，结果已验证”。"),
    ]
    for table, number, body in demo_rows:
        cells = unique_cells(table.rows[0])
        set_cell_text(cells[0], number, size=7.0, bold=True, fill=BLUE, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[-1], body, size=6.9, fill=WHITE)
    demo_caption = find_paragraph(document, "高价值闭环：发现问题 → 给方案 → 人确认 → 安全执行 → 证明结果")
    # Remove the small redundant diagram immediately above the caption to make room for fallback evidence.
    previous = demo_caption._p.getprevious()
    if previous is not None and previous.tag.endswith("}p"):
        previous.getparent().remove(previous)
    add_demo_fallback(document, demo_caption._p)

    # Current value page stays evidence-safe.
    value_table = tables[64]
    replace_visual_row(value_table, "储能管理", ["储能管理", "SOC / 备电目标、充放电计划执行情况", "Simulator E2E 已验证；真实 Anker 实机待补"], size=7.2)
    replace_visual_row(value_table, "峰值管理", ["峰值管理", "家庭峰值功率变化", "Simulator 场景可做执行前后 Smart Meter 对比"], size=7.2)
    replace_visual_row(value_table, "减少重复操作", ["减少重复操作", "一次家庭目标协调多个设备角色", "4 项 Action Set E2E 已验证"], size=7.2)
    replace_visual_row(value_table, "不夸大长期效果", ["不夸大长期效果", "长期节省 / 策略价值", "PENDING_LONG_TERM；示例仅说明计算模型"], size=7.2)

    # Add quantitative value model and SOLIX ecosystem value after the value page.
    innovation_anchor = find_paragraph(document, "04 创新点")._p
    add_value_solix_page(document, innovation_anchor, divider_template)

    # Reorder innovation points, L3 first.
    innovation_texts = [
        (tables[66], "01", "从“设备执行成功”升级为“家庭结果被证明”｜L1 命令确认 + L2 设备确认 + L3 Smart Meter 家庭结果确认；只有家庭结果符合预期，方案才最终 VERIFIED。"),
        (tables[67], "02", "一个家庭目标，协调多个设备角色｜用户只表达“减少光伏反送”“多留一点电”；系统组织一次授权、顺序执行、逐步回读、异常 Fail-stop 的多设备行动方案。"),
        (tables[68], "03", "AI 智能与真实设备执行权限分离｜Hermes 能理解和解释，但不能绕过用户授权、执行前安全检查与系统控制中枢直接写设备。"),
        (tables[69], "04", "确定性优化 + AI 理解明确分工｜AI 负责语言、目标、解释与澄清；Optimizer 负责具体数值、设备边界、硬约束与可重复结果。"),
        (tables[70], "05", "从 Dashboard 到 Autonomous Home Energy Agent｜系统持续观察、预测与主动判断；无需动作时保持安静，有价值时才主动介入。"),
    ]
    for table, number, body in innovation_texts:
        cells = unique_cells(table.rows[0])
        set_cell_text(cells[0], number, size=8.0, bold=True, fill=BLUE, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[-1], body, size=7.6, fill=WHITE)

    # Engineering evidence and status truth.
    engineering = tables[74]
    replace_visual_row(engineering, "真实 Anker SOLIX", ["SOLIX-compatible 写入 / 回读", "Home Assistant + TCP Simulator 写入与状态回读", "SIMULATOR_E2E｜VERIFIED"], size=6.9)
    replace_visual_row(engineering, "真实跨设备 Action Set", ["4 项跨设备 Action Set", "一次授权、顺序执行、每步状态回读、异常 Fail-stop", "SIMULATOR_E2E｜VERIFIED"], size=6.9)
    replace_visual_row(engineering, "Smart Meter 家庭结果验证", ["Smart Meter L3 家庭结果确认", "记录执行前后家庭电网变化并验证方向", "SIMULATOR_E2E｜VERIFIED"], size=6.9)
    replace_visual_row(engineering, "Voice", ["Voice 增强交互", "软件 / 服务链路已实现；ESP32-S3 实板语音 HIL 待验收", "SOFTWARE_IMPLEMENTED｜HIL_PENDING"], size=6.7)
    replace_visual_row(engineering, "Simulator", ["真实 Anker 家庭硬件", "写入、长期节省、准确率与用户规模", "NOT_VERIFIED"], size=6.9)

    # Final page: lock the three evaluator takeaways.
    final_table = tables[78]
    cell = unique_cells(final_table.rows[0])[0]
    set_cell_text(
        cell,
        "评委只需记住三句话\n\n01｜光衡会主动判断，不是等用户操作。\n02｜一个家庭目标可以协调多个设备角色，但 AI 不能越过用户授权边界。\n03｜命令成功不算结束，只有 Smart Meter 证明家庭结果真的改变，任务才 VERIFIED。",
        size=9.0,
        bold=True,
        color=NAVY,
        fill=PALE_GREEN,
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )
    set_paragraph_text(
        document.paragraphs[-1],
        "让“家庭能源智能”从会说，走到会做；从会做，走到有证据。",
        size=10.0,
        bold=True,
        color=NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
