from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

from docx import Document


SOURCE = Path(r"C:\Users\bandu\Desktop\黑客松比赛材料\光衡_GuangHeng_Anker黑客松预赛提交材料_20页优化版(1).docx")
FINAL = Path(r"C:\Users\bandu\Desktop\黑客松比赛材料\光衡_GuangHeng_Anker黑客松预赛提交材料_最终强化版.docx")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def all_text(document: Document) -> str:
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                marker = id(cell._tc)
                if marker not in seen:
                    seen.add(marker)
                    blocks.append(cell.text)
    return "\n".join(blocks)


def main() -> None:
    with zipfile.ZipFile(FINAL) as archive:
        bad = archive.testzip()
        if bad:
            raise RuntimeError(f"Corrupt ZIP member: {bad}")

    document = Document(FINAL)
    text = all_text(document)
    required = [
        "命令成功  ≠  设备成功  ≠  家庭目标成功",
        "L3 Smart Meter 家庭结果确认",
        "AI 技术路径：让 AI 负责理解，让确定性系统负责控制",
        "新增自用电量 ＝ 年光伏发电量",
        "年净收益 ≈ 新增自用电量",
        "示例测算｜仅用于说明模型，不代表当前实测收益",
        "对 Anker SOLIX 的价值",
        "Demo Reliability｜现场容灾",
        "SOFTWARE_IMPLEMENTED｜HIL_PENDING",
        "PENDING_LONG_TERM",
        "[二维码 / 视频链接待补]",
        "评委只需记住三句话",
    ]
    forbidden = [
        "五个真实用户痛点",
        "REAL_HIL｜VERIFIED真实跨设备 Action Set",
        "OpenAI GPT",
        "MPC",
        "强化学习",
        "年省 ¥",
    ]
    missing = [item for item in required if item not in text]
    found_forbidden = [item for item in forbidden if item in text]
    print(f"source_sha256={sha256(SOURCE)}")
    print(f"final_sha256={sha256(FINAL)}")
    print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} images={len(document.inline_shapes)}")
    print(f"missing_required={missing}")
    print(f"found_forbidden={found_forbidden}")
    if sha256(SOURCE) != "85F2F47F5D8BFFD544207C3F7C302A75DC5EE6AF6CA69ABE29965FF35D065C34":
        raise RuntimeError("Source file hash changed")
    if missing or found_forbidden:
        raise RuntimeError("Content QA failed")
    print("QA=PASS")


if __name__ == "__main__":
    main()
